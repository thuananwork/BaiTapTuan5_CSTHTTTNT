"""Tao bao cao Word 8 muc tu cac tep trong results/ (so lieu lay thang tu CSV/JSON)."""
import json, pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TAB = os.path.join(ROOT, "results", "tables") + os.sep
FIGD = os.path.join(ROOT, "results", "figures") + os.sep
OUT = os.path.join(ROOT, "report", "BaoCao_TabICLv2_BoHoc.docx")
info = json.load(open(TAB + "thong_tin_chung.json"))
ex = json.load(open(TAB + "vi_du_attention.json"))
cond = json.load(open(TAB + "so_dieu_kien.json"))
cfg = pd.read_csv(TAB + "cau_hinh_logistic.csv")
opt = pd.read_csv(TAB + "toi_uu_toc_do_hoc.csv")
thr = pd.read_csv(TAB + "nguong_quyet_dinh.csv")
att = pd.read_csv(TAB + "attention_ngu_canh.csv")
cmpw = pd.read_csv(TAB + "doi_chieu_tu_cai_sklearn.csv")
dq = pd.read_csv(TAB + "thong_ke_dinh_luong.csv", index_col=0)
dt = pd.read_csv(TAB + "thong_ke_dinh_tinh.csv")
corr = pd.read_csv(TAB + "tuong_quan.csv", index_col=0)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = sec.bottom_margin = Cm(2.3)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.25
for name, size in [("Heading 1", 16), ("Heading 2", 13.5), ("Heading 3", 12)]:
    h = doc.styles[name]; h.font.name = "Times New Roman"; h.font.size = Pt(size)
    h.font.bold = True; h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    h.paragraph_format.space_after = Pt(6)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color); tcPr.append(s)


def P(text, bold=False, italic=False, align=None, size=None, after=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align != "l": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p


def RICH(parts):
    """parts: list of (text, bold)"""
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for t, b in parts:
        r = p.add_run(t); r.bold = b
    return p


def F(text):  # cong thuc: phu tro, can giua, nghieng
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.italic = True; r.font.name = "Cambria Math"; r.font.size = Pt(12)
    p.paragraph_format.space_after = Pt(6)


def H1(t): doc.add_heading(t, 1)
def H2(t): doc.add_heading(t, 2)


def BUL(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(3)


def TABLE(header, rows, widths=None, size=10, hl=None):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(size)
        shade(c, "DCE6F1")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; para = cells[i].paragraphs[0]; r = para.add_run(str(v)); r.font.size = Pt(size)
            if i > 0 and isinstance(v, str) and v.replace(".", "").replace(",", "").replace("%", "").replace("−", "").replace("-", "").isdigit():
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if hl is not None and ri in hl: shade(cells[i], "FFF2CC")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def FIG(fn, cap, w=15.0):
    doc.add_picture(FIGD + fn, width=Cm(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(cap, italic=True, align="c", size=10.5, after=10)


def BOX(title, text, fill="FFF2CC"):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"; c = t.rows[0].cells[0]; shade(c, fill)
    c.text = ""; p = c.paragraphs[0]; r = p.add_run(title + " "); r.bold = True
    p.add_run(text); p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def th(n): return format(int(round(n)), ",").replace(",", ".")


def f(x, k=3): return f"{x:.{k}f}".replace(".", ",")
def pc(x, k=1): return f"{100 * x:.{k}f}%".replace(".", ",")


# Điền tên vào đây rồi chạy lại script để bìa có tên (hoặc sửa thẳng trong Word)
THANH_VIEN = ".............................................................."
GIANG_VIEN = ".............................................................."

# ======================= TRANG BIA =======================
for _ in range(3): P("", after=12)
P("BÁO CÁO HỌC PHẦN", bold=True, align="c", size=14)
P("CƠ SỞ TOÁN HỌC CHO TRÍ TUỆ NHÂN TẠO", bold=True, align="c", size=18)
P("", after=18)
P("Nghiên cứu cơ sở toán học của mô hình TabICLv2 và ứng dụng dự đoán nguy cơ sinh viên bỏ học",
  bold=True, align="c", size=16)
P("", after=24)
P("Lĩnh vực: Giáo dục", align="c"); P("Bộ dữ liệu: Predict Students' Dropout and Academic Success (UCI), 4.424 sinh viên", align="c")
P("", after=24)
P("Nhóm thực hiện: Nhóm 6", align="c")
P("Thành viên: " + THANH_VIEN, align="c")
P("Giảng viên: " + GIANG_VIEN, align="c")
P("", after=12); P("Tháng 10 năm 2026", align="c")
doc.add_page_break()

# ======================= TOM TAT =======================
H1("Tóm tắt")
P("Bài báo cáo đi từ một câu hỏi rất đời thường của nhà trường: sinh viên nào đang đứng trước nguy cơ bỏ học, "
  "và có thể biết điều đó đủ sớm để giúp được hay không. Nhà trường cần đoán trước em nào có nguy cơ bỏ học để kịp hỗ trợ. "
  "Ở đây mỗi sinh viên là một mẫu dữ liệu; tốt nghiệp là nhóm an toàn (nhãn 0), bỏ học là nhóm nguy cơ (nhãn 1).")
P(f"Dữ liệu gồm {th(info['n_raw'])} sinh viên của Viện Bách khoa Portalegre (Bồ Đào Nha). Sau khi bỏ nhóm còn đang học, "
  f"còn {th(info['n_used'])} sinh viên với tỉ lệ bỏ học {pc(info['dropout_rate_used'])}. Mô hình được chọn là TabICLv2, một mô hình nền "
  "cho dữ liệu dạng bảng, học trong ngữ cảnh. Để thấy rõ từng mảnh toán, báo cáo dùng thêm Hồi quy Logistic tự cài bằng "
  "Gradient Descent làm mô hình đối chứng.")
BOX("Điều cần nói thẳng:",
    "nhóm đã chạy TabICLv2 thật (thư viện tabicl 2.2.0, trọng số tabicl-classifier-v2-20260212, 8 thành viên ensemble, toàn bộ tập huấn luyện làm ngữ cảnh) "
    "qua notebook 05, trên đúng cách chia dữ liệu và bốn cấu hình như mô hình đối chứng. Kết quả ở mục 7.5. "
    "Phần chưa làm: khảo sát kích thước ngữ cảnh và số thành viên ensemble, kiểm tra chéo. "
    "Thí nghiệm về tốc độ học vẫn làm trên Hồi quy Logistic vì TabICLv2 không được huấn luyện lại trên dữ liệu của nhóm.")

# ======================= 1 =======================
H1("1. Phát biểu và mô hình hóa bài toán")
BOX("Nói dễ hiểu:", 'Trước khi tính gì, ta phải nói rõ mình đang đoán điều gì. Ở đây là: nhìn vài con số của một sinh viên, đoán em đó sau này tốt nghiệp hay bỏ học. Ta muốn biết điều đó trước khi em ấy nghỉ học, khi vẫn còn kịp giúp. Mọi công thức phía sau chỉ để trả lời đúng câu hỏi đơn giản này.', fill="E2F0D9")
H2("1.1 Bối cảnh và mục tiêu")
P("Sinh viên bỏ học là mất mát cho cả hai phía: nhà trường mất người học, còn sinh viên mất thời gian và tiền bạc. "
  "Khi đơn thôi học đã nộp thì không còn gì để cứu. Điều nhà trường cần là một cảnh báo đến sớm, khi cố vấn học tập "
  "vẫn còn kịp nhắn tin, gặp mặt hoặc hỗ trợ học phí. Mục tiêu của bài toán vì vậy là ước lượng, từ những thông tin "
  "đã biết về một sinh viên, xác suất em đó sẽ bỏ học.")
H2("1.2 Nhãn an toàn và nhãn nguy cơ")
P("Muốn có bài toán thì phải nói rõ thế nào là thành công, thế nào là thất bại. Mô hình nhìn các chỉ số của một sinh viên để đoán em đó có bỏ học không, "
  "trước khi chuyện xảy ra. Ta quy ước sinh viên tốt nghiệp (Graduate) là nhóm an toàn, nhãn y = 0; sinh viên bỏ học (Dropout) là nhóm nguy cơ, nhãn y = 1.")
P(f"Nhóm sinh viên đang học (Enrolled, {info['raw_counts']['Enrolled']} người) là nhóm chưa có kết quả: chưa ai biết kết cục của họ. "
  "Đưa họ vào thì mô hình sẽ học từ một nhãn chưa có kết luận, nên nhóm loại họ khỏi bài toán. Còn lại "
  f"{th(info['n_used'])} sinh viên, trong đó {pc(info['dropout_rate_used'])} đã bỏ học.")
H2("1.3 Loại bài toán, đầu vào và đầu ra")
P("Đây là bài toán phân loại nhị phân. Đầu ra của mô hình không chỉ là “bỏ học” hay “không”, mà là một xác suất "
  "p = P(y = 1 | x) nằm giữa 0 và 1, để nhà trường tự chọn mức báo động. Đầu vào là một số ít chỉ số có lý do giáo dục rõ ràng. "
  "Bộ dữ liệu có 36 cột, nhưng để trả lời câu hỏi ở trên chỉ cần tám cột dưới đây; các cột còn lại (nghề nghiệp bố mẹ, GDP, lạm phát...) "
  "tạm để ngoài vì không cần cho câu hỏi gốc. Đề bài cho phép nêu lý do không dùng một biến, và đây là lý do của nhóm.")
TABLE(["Ký hiệu", "Tên cột trong dữ liệu", "Ý nghĩa", "Kiểu"],
      [["x₁", "Tuition fees up to date", "Đóng học phí đủ hạn (1 = có)", "Nhị phân"],
       ["x₂", "Debtor", "Đang nợ (1 = có)", "Nhị phân"],
       ["x₃", "Scholarship holder", "Có học bổng (1 = có)", "Nhị phân"],
       ["x₄", "Age at enrollment", "Tuổi lúc nhập học", "Số nguyên"],
       ["x₅", "1st sem (approved)", "Số môn qua ở học kỳ 1", "Số nguyên"],
       ["x₆", "1st sem (grade)", "Điểm trung bình học kỳ 1 (thang 20)", "Số thực"],
       ["x₇", "2nd sem (approved)", "Số môn qua ở học kỳ 2", "Số nguyên"],
       ["x₈", "2nd sem (grade)", "Điểm trung bình học kỳ 2 (thang 20)", "Số thực"]],
      widths=[1.6, 4.6, 7.2, 2.6])
H2("1.4 Biểu diễn bằng toán học")
P("Mỗi sinh viên là một vector gồm tám số, và cả tập dữ liệu là một ma trận mà mỗi hàng là một sinh viên. "
  "Nhãn của cả tập xếp thành một vector cột.")
F("xᵢ ∈ ℝ⁸,    X ∈ ℝⁿˣ⁸ (n = 3.630),    y ∈ {0, 1}ⁿ")
P("Việc cần học là một hàm f nhận vector x và trả về xác suất bỏ học. Hàm đó có các tham số θ (với Hồi quy Logistic "
  "là trọng số w và hệ số chặn b; với TabICLv2 là một số lượng rất lớn trọng số của mạng Transformer). Bài toán tối ưu là "
  "tìm θ làm cho mức “đoán sai” trên dữ liệu huấn luyện nhỏ nhất, với một ràng buộc duy nhất: đầu ra phải nằm trong [0, 1] "
  "để đọc được như xác suất.")
F("θ* = argmin_θ  L(θ) = argmin_θ  (1/n) Σᵢ ℓ( f_θ(xᵢ), yᵢ ),    0 ≤ f_θ(x) ≤ 1")
H2("1.5 Hai mô hình trong báo cáo")
P("Mô hình chính của đề tài là TabICLv2 (Qu, Holzmüller, Varoquaux, Le Morvan, ICML 2026), một mô hình nền cho dữ liệu dạng bảng. "
  "Điểm đặc biệt của nó là không huấn luyện lại trên dữ liệu của mình: nó nhìn các sinh viên đã biết kết cục như “ngữ cảnh”, "
  "rồi dự đoán cho sinh viên mới chỉ bằng một lượt tính. Vì TabICLv2 khá phức tạp, nhóm dùng thêm Hồi quy Logistic làm mô hình đối chứng: "
  "nó đơn giản đến mức mọi mảnh toán (vector, tích vô hướng, đạo hàm, Gradient Descent) đều nhìn thấy được, "
  "và là nơi nhóm làm các thí nghiệm tối ưu mà đề bài yêu cầu.")

# ======================= 2 =======================
H1("2. Thống kê mô tả và đặc điểm dữ liệu")
BOX("Nói dễ hiểu:", 'Trước khi dạy máy, hãy tự nhìn dữ liệu như người làm khảo sát: có bao nhiêu em, bao nhiêu em bỏ học, nhóm nào bỏ học nhiều. Việc này cho ta biết nên tin vào chỉ số nào, và nên cẩn thận ở đâu (ví dụ điểm bằng 0 thật ra có nghĩa gì).', fill="E2F0D9")
H2("2.1 Giới thiệu bộ dữ liệu")
P(f"Bộ dữ liệu Predict Students' Dropout and Academic Success do Realinho, Martins, Machado và Baptista công bố (UCI Machine Learning Repository, "
  f"DOI 10.24432/C5MC89; bài mô tả trên tạp chí Data, 2022). Dữ liệu gồm sinh viên của 17 ngành tại Viện Bách khoa Portalegre, "
  f"Bồ Đào Nha, nhập học từ 2008 đến 2019. Có {th(info['n_raw'])} sinh viên, {info['n_features']} đặc trưng và một biến mục tiêu ba lớp. "
  f"Dữ liệu không có giá trị thiếu (đã kiểm tra: {info['missing']} ô trống).")
H2("2.2 Biến định lượng")
rows = []
nm = {"Age at enrollment": "Tuổi nhập học", "Curricular units 1st sem (approved)": "Số môn qua HK1",
      "Curricular units 1st sem (grade)": "Điểm HK1", "Curricular units 2nd sem (approved)": "Số môn qua HK2",
      "Curricular units 2nd sem (grade)": "Điểm HK2"}
for k, v in dq.iterrows():
    rows.append([nm[k], f(v["mean"], 2), f(v["median"], 2), f(v["std"], 2), f(v["variance"], 1), f(v["min"], 1),
                 f(v["25%"], 1), f(v["75%"], 1), f(v["max"], 1)])
TABLE(["Biến", "Trung bình", "Trung vị", "Độ lệch chuẩn", "Phương sai", "Min", "Q1", "Q3", "Max"], rows, size=9)
P(f"Điểm HK1 có trung bình {f(dq.loc['Curricular units 1st sem (grade)','mean'],2)} nhưng trung vị tới {f(dq.loc['Curricular units 1st sem (grade)','median'],2)}. "
  "Trung bình thấp hơn trung vị cho thấy phân bố lệch trái, bị kéo xuống bởi một nhóm sinh viên có điểm bằng 0. "
  "Tuổi nhập học thì ngược lại: trung vị 20 nhưng trung bình 23,5 và lớn nhất 70, tức có một đuôi sinh viên lớn tuổi kéo trung bình lên.")
H2("2.3 Biến định tính")
rows = []
vi = {"Chua dong du": "Chưa đóng đủ học phí", "Da dong du": "Đã đóng đủ học phí", "Khong no": "Không nợ", "Dang no": "Đang nợ",
      "Khong hoc bong": "Không có học bổng", "Co hoc bong": "Có học bổng", "Nu": "Nữ", "Nam": "Nam"}
for _, r in dt.iterrows():
    rows.append([vi[r.nhom], int(r.so_sv), pc(r.ty_le), pc(r.ty_le_bo_hoc)])
TABLE(["Nhóm", "Số sinh viên", "Tỉ lệ trong dữ liệu", "Tỉ lệ bỏ học"], rows, widths=[6, 3, 4, 3.5])
P("Bảng cho thấy những khác biệt rất lớn. Trong nhóm chưa đóng đủ học phí, 94,0% bỏ học, so với 30,7% ở nhóm đã đóng đủ. "
  "Sinh viên đang nợ bỏ học 75,5%, còn không nợ là 34,5%. Có học bổng thì chỉ 13,8% bỏ học, không có thì 48,4%. "
  "Nam bỏ học nhiều hơn nữ (56,1% so với 30,2%). Đây đều là những chênh lệch mô hình có thể khai thác, và chúng khớp với trực giác: "
  "tiền bạc là một trong những lý do thực tế nhất khiến sinh viên rời trường. Giới tính là một yếu tố cần cẩn trọng khi dùng "
  "ngoài đời, vì dự đoán không nên biến thành định kiến; nhóm không đưa nó vào bộ tám cột chính.")
FIG("h4_ty_le_theo_nhom.png", "Hình 1. Tỉ lệ bỏ học theo từng nhóm. Đường đứt là tỉ lệ chung (39,1%).", 11.5)
H2("2.4 Biến mục tiêu và mất cân bằng")
P(f"Nhãn gốc gồm Graduate {th(info['raw_counts']['Graduate'])}, Dropout {th(info['raw_counts']['Dropout'])} và Enrolled {info['raw_counts']['Enrolled']}. "
  f"Sau khi bỏ Enrolled, lớp bỏ học chiếm {pc(info['dropout_rate_used'])}. Mức này lệch nhưng chưa nặng: nếu mô hình ngây thơ luôn đoán “an toàn” "
  f"thì vẫn đúng khoảng {pc(1 - info['dropout_rate_used'])} nhưng không phát hiện được em bỏ học nào. Vì thế Accuracy một mình không đủ để đánh giá, "
  "và báo cáo luôn xem thêm Recall của lớp bỏ học. Việc chia train/test cũng phải giữ nguyên tỉ lệ này (chia phân tầng, mục 6.2).")
FIG("h1_phan_bo_nhan.png", "Hình 2. Phân bố nhãn trước và sau khi bỏ nhóm đang học.", 13)
H2("2.5 Dữ liệu thiếu, ngoại lệ và thang đo")
P(f"Dữ liệu không thiếu. Ngoại lệ theo quy tắc IQR (ngoài khoảng Q1 − 1,5·IQR đến Q3 + 1,5·IQR) có ở mọi biến: tuổi {info['outliers_iqr']['Age at enrollment']} sinh viên, "
  f"điểm HK1 {info['outliers_iqr']['Curricular units 1st sem (grade)']}, điểm HK2 {info['outliers_iqr']['Curricular units 2nd sem (grade)']}. "
  f"Con số lớn ở điểm số có nguyên nhân rõ: {info['zero_grade_sem2']} sinh viên có điểm HK2 bằng 0, và trong nhóm đó "
  f"{pc(info['zero_grade_sem2_dropout_rate'])} đã bỏ học. Đây không phải lỗi nhập liệu. Điểm 0 thường nghĩa là em đó không còn đi thi, "
  "tức chính là tín hiệu của việc rời trường, nên nhóm giữ nguyên các giá trị này.")
P("Nhưng điều đó kéo theo một lưu ý quan trọng sẽ được nhắc lại ở phần hạn chế: dữ liệu học kỳ 2 phần nào đã chứa sẵn đáp án, "
  "vì khi một em đã nghỉ giữa năm thì điểm học kỳ 2 của em là 0. Cảnh báo dựa vào học kỳ 2 vì vậy đến muộn hơn cảnh báo dựa vào học kỳ 1.")
P("Về thang đo, các cột chênh lệch nhau rất nhiều: cột nhị phân chỉ có 0 và 1, điểm số từ 0 đến khoảng 19, còn tuổi từ 17 đến 70. "
  "Sự chênh lệch này không chỉ là chuyện thẩm mỹ; nó ảnh hưởng trực tiếp đến tốc độ hội tụ của Gradient Descent và sẽ được đo ở mục 5.")
FIG("h2_boxplot.png", "Hình 3. Số môn qua và tuổi nhập học theo nhóm. Nhóm bỏ học (phải) có số môn qua thấp hơn hẳn.", 15)
H2("2.6 Mối quan hệ giữa các biến")
a, b = "Curricular units 1st sem (approved)", "Curricular units 2nd sem (approved)"
P(f"Tương quan với nhãn bỏ học mạnh nhất là số môn qua ở HK2 ({f(corr.loc[b,'y'],2)}), rồi điểm HK2 ({f(corr.loc['Curricular units 2nd sem (grade)','y'],2)}), "
  f"số môn qua HK1 ({f(corr.loc[a,'y'],2)}) và đóng học phí ({f(corr.loc['Tuition fees up to date','y'],2)}). Dấu âm nghĩa là qua nhiều môn, đóng đủ học phí thì ít bỏ học. "
  f"Điều đáng chú ý hơn nằm giữa các đặc trưng với nhau: số môn qua HK1 và HK2 có tương quan tới {f(corr.loc[a,b],2)}. "
  "Hai cột này gần như kể cùng một câu chuyện, và hiện tượng cộng tuyến này sẽ giải thích một chi tiết lạ trong trọng số của mô hình ở mục 4.4.")
FIG("h3_tuong_quan.png", "Hình 4. Ma trận tương quan của tám đặc trưng và nhãn bỏ học.", 11.5)

# ======================= 3 =======================
H1("3. Cơ sở đại số tuyến tính trong biểu diễn và xử lý dữ liệu")
BOX("Nói dễ hiểu:", 'Máy tính chỉ hiểu con số. Nên ta viết mỗi sinh viên thành một dãy 8 số (vector), xếp cả lớp thành một bảng số (ma trận), và gộp 8 số lại thành một “mức nguy hiểm” bằng cách nhân từng số với một trọng số rồi cộng lại (tích vô hướng). Phần attention của TabICLv2 cũng chỉ là cách so sánh các dãy số bằng nhân và cộng như vậy.', fill="E2F0D9")
H2("3.1 Sinh viên là một vector, cả lớp là một ma trận")
P("Máy không hiểu câu “em An hay nghỉ học và nợ học phí”; nó chỉ hiểu số. Vì vậy mỗi sinh viên được viết thành một dãy tám số theo đúng thứ tự "
  "ở bảng mục 1.3, tức một vector trong không gian tám chiều. Ví dụ tự lập cho minh họa: x = (1, 0, 0, 19, 5, 12,3, 5, 12,1) là một sinh viên "
  "đóng đủ học phí, không nợ, không học bổng, 19 tuổi, qua 5 môn mỗi kỳ với điểm khoảng 12. Xếp 3.630 vector này chồng lên nhau ta được ma trận X "
  "kích thước 3.630 × 8. Mọi tính toán về sau, từ dự đoán cho đến tính gradient, đều là phép nhân ma trận trên X, nên máy tính xử lý rất nhanh.")
H2("3.2 Tích vô hướng: gộp các chỉ số thành một “mức nguy hiểm”")
P("Các chỉ số không nặng nhẹ như nhau: nợ học phí đáng lo hơn một chút chênh lệch về tuổi. Cách tự nhiên nhất để gộp chúng là nhân mỗi chỉ số "
  "với một trọng số rồi cộng lại. Phép “nhân từng cặp rồi cộng” chính là tích vô hướng giữa vector trọng số w và vector đặc trưng x. "
  "Trọng số lớn nghĩa là chỉ số đó có ảnh hưởng mạnh, dấu của trọng số cho biết nó đẩy nguy cơ lên hay xuống.")
F("z = wᵀx + b = w₁x₁ + w₂x₂ + … + w₈x₈ + b")
P("Cũng nhờ tích vô hướng, ta tính được cho cả lớp trong một lần: z = Xw + b là vector gồm mức nguy hiểm của mọi sinh viên.")
H2("3.3 Chuẩn hóa, khoảng cách và độ tương đồng")
P("Để so sánh hai sinh viên, ta cần một thước đo “giống nhau”. Cách đơn giản là khoảng cách Euclid, tức độ dài (chuẩn) của hiệu hai vector. "
  "Nhưng nếu để nguyên thang đo thì tuổi (17 đến 70) sẽ lấn át cột nhị phân (0 hoặc 1) chỉ vì số của nó lớn hơn, chứ không vì nó quan trọng hơn. "
  "Vì vậy trước hết mỗi cột được đưa về cùng thang bằng phép chuẩn hóa z-score: trừ trung bình và chia độ lệch chuẩn, trong đó trung bình "
  "và độ lệch chuẩn tính trên tập huấn luyện.")
F("x̃ⱼ = (xⱼ − μⱼ) / σⱼ,        d(x, x′) = ‖x̃ − x̃′‖₂ = √Σⱼ (x̃ⱼ − x̃′ⱼ)²")
H2("3.4 Các phép toán trong TabICLv2")
P("TabICLv2 là một mô hình Transformer, và lõi của Transformer là attention, được viết hoàn toàn bằng phép nhân ma trận. "
  "Mỗi sinh viên đã được biến thành một vector biểu diễn; nhân vector đó với ba ma trận trọng số học được sẽ cho ra ba vector gọi là truy vấn (Q), khóa (K) và giá trị (V). "
  "Mức độ “chú ý” của sinh viên cần dự đoán vào một sinh viên trong ngữ cảnh là tích vô hướng giữa truy vấn của người này với khóa của người kia, "
  "rồi qua softmax để thành các trọng số cộng lại bằng 1.")
F("Q = XW_Q,  K = XW_K,  V = XW_V,      Attention(Q, K, V) = softmax( QKᵀ / √d ) · V")
P("Theo bài gốc, TabICLv2 gồm ba tầng nối nhau: một tầng nhúng theo từng cột để biến mỗi đặc trưng thành vector, một tầng tương tác theo hàng "
  "gộp các đặc trưng của một sinh viên vào một vector duy nhất, và tầng cuối là học trong ngữ cảnh theo toàn bộ tập dữ liệu, nơi các sinh viên cần dự đoán "
  "“chú ý” vào các sinh viên đã có nhãn. Nhãn của sinh viên huấn luyện được cộng thêm vào biểu diễn của họ, nhờ đó thông tin kết cục đi theo trọng số attention "
  "sang người cần dự đoán. (Mô tả này tóm tắt từ bài gốc, cần đối chiếu chi tiết khi nộp bản cuối.)")
H2("3.5 Ví dụ tính tay: một đầu attention với bốn sinh viên")
P("Để thấy attention thực sự làm gì, nhóm tính tay trên một ví dụ nhỏ. Chỉ dùng hai đặc trưng (số môn qua HK1 và HK2), bốn sinh viên trong ngữ cảnh "
  "lấy từ tập huấn luyện và một sinh viên cần dự đoán lấy từ tập test. Đây là bản đơn giản hóa để minh họa cơ chế, không có trọng số học được "
  "nên không phải TabICLv2: độ giống nhau lấy thẳng từ khoảng cách, điểm attention là −d²/2 (τ = 1).")
rows = []
for i in range(4):
    rows.append([f"Sinh viên {i+1}", f"({int(ex['raw_ctx'][i][0])}; {int(ex['raw_ctx'][i][1])})",
                 f"({f(ex['z_ctx'][i][0])}; {f(ex['z_ctx'][i][1])})", "Bỏ học" if ex['nhan_ctx'][i] else "An toàn",
                 f(ex['diem'][i]), f(ex['trong_so'][i])])
TABLE(["Ngữ cảnh", "Số môn qua (HK1; HK2)", "Sau chuẩn hóa", "Kết cục thật", "Điểm −d²/2", "Trọng số attention"], rows, size=9.5)
P(f"Sinh viên cần dự đoán qua (1; 1) môn, sau chuẩn hóa là ({f(ex['z_q'][0])}; {f(ex['z_q'][1])}). Trung bình và độ lệch chuẩn dùng để chuẩn hóa "
  f"là μ = ({f(ex['mu'][0],2)}; {f(ex['mu'][1],2)}), σ = ({f(ex['sd'][0],2)}; {f(ex['sd'][1],2)}). Ví dụ với sinh viên 1: khoảng cách bình phương tới người cần dự đoán là "
  f"(−1,502 + 1,191)² + (−1,446 + 1,128)² ≈ 0,198, nên điểm là −0,198/2 ≈ −0,099. Softmax biến bốn điểm thành bốn trọng số cộng lại bằng 1. "
  f"Hai sinh viên có hồ sơ gần nhất, cả hai đều bỏ học, nhận gần hết trọng số (0,422 và 0,444). Xác suất bỏ học dự đoán là tổng trọng số của những người bỏ học: "
  f"0,422 + 0,444 ≈ 0,866 (với số chưa làm tròn là {f(ex['p_bo_hoc'],3)}). Kết cục thật của sinh viên này là bỏ học, nên dự đoán đúng.")
H2("3.6 Trị riêng và những gì không dùng trực tiếp")
P("Trị riêng xuất hiện ở bài này qua ma trận XᵀX/n. Trị riêng lớn nhất cho biết hàm mất mát “dốc” nhất theo hướng nào, và tỉ số giữa trị riêng lớn nhất "
  f"và nhỏ nhất (số điều kiện) cho biết bề mặt đó méo đến đâu. Khi chuẩn hóa dữ liệu, số điều kiện là khoảng {f(cond['chuan_hoa']['so_dieu_kien'],1)}; khi để nguyên thang đo, nó vọt lên khoảng "
  f"{th(cond['khong_chuan_hoa']['so_dieu_kien'])}" + ". Đây là cơ sở toán học để giải thích thí nghiệm ở mục 5.4. "
  "Còn các phép phân rã như SVD hay PCA thì bài này không dùng trực tiếp: Hồi quy Logistic và TabICLv2 đều không đòi giảm chiều, và chỉ tám đặc trưng thì cũng không cần giảm.")

# ======================= 4 =======================
H1("4. Cơ sở giải tích và xây dựng hàm mất mát")
BOX("Nói dễ hiểu:", 'Có mức nguy hiểm rồi, ta cần hai thứ: một cách đổi nó thành xác suất (sigmoid), và một thước đo xem đoán sai nặng đến đâu (hàm mất mát). Đoán sai mà còn tự tin thì phải bị phạt rất nặng, chẳng hạn khẳng định em an toàn mà em lại bỏ học. Đạo hàm chỉ cho ta biết nên chỉnh các trọng số theo hướng nào để mức phạt giảm.', fill="E2F0D9")
H2("4.1 Mô hình: từ mức nguy hiểm đến xác suất")
P("Mức nguy hiểm z có thể là −5 hay 37, khó đọc. Giáo viên muốn nghe “em này có 80% khả năng bỏ học”. Ta cần một hàm nén mọi số thực về khoảng (0, 1), "
  "và hàm sigmoid làm đúng việc đó: z rất âm cho xác suất gần 0, z bằng 0 cho đúng 0,5, z rất dương cho xác suất gần 1. "
  "Kết quả là xác suất có điều kiện: biết các chỉ số của em này, khả năng em bỏ học là bao nhiêu.")
F("p = P(y = 1 | x) = σ(z) = 1 / (1 + e^(−z)),        σ′(z) = σ(z)(1 − σ(z))")
H2("4.2 Hàm mất mát cross-entropy và lý do chọn")
P("Ta cần một con số đo xem mô hình đoán sai đến mức nào. Ví dụ: nếu mô hình tự tin nói “em này an toàn 99%” mà thực ra em bỏ học, "
  "nó phải bị phạt rất nặng; nếu chỉ lưỡng lự 50/50 thì phạt vừa phải. Hàm cross-entropy làm đúng điều đó nhờ logarit. Với một sinh viên, mức phạt là "
  "−ln p nếu em thật sự bỏ học, và −ln(1 − p) nếu em thật sự an toàn.")
F("L(w, b) = −(1/n) Σᵢ [ yᵢ ln pᵢ + (1 − yᵢ) ln(1 − pᵢ) ]")
TABLE(["Mô hình dự đoán P(bỏ học)", "Thật sự bỏ học: mức phạt −ln p", "Thật sự an toàn: mức phạt −ln(1 − p)"],
      [["0,99 (rất tự tin là bỏ học)", "0,010", "4,605"], ["0,90", "0,105", "2,303"], ["0,50 (lưỡng lự)", "0,693", "0,693"],
       ["0,10", "2,303", "0,105"], ["0,01 (rất tự tin là an toàn)", "4,605", "0,010"]], widths=[5.5, 5.5, 5.5])
P("Bảng cho thấy mức phạt tăng vọt khi mô hình tự tin mà sai (từ 0,010 lên 4,605), đúng như tinh thần “nổ tan xác”. "
  "Lý do sâu hơn nằm ở xác suất: nếu coi mỗi sinh viên là một lần tung đồng xu lệch với xác suất bỏ học pᵢ, thì khả năng quan sát đúng bộ nhãn y của cả tập "
  "là tích của các pᵢ hoặc (1 − pᵢ). Lấy logarit rồi đổi dấu và chia n thì được đúng hàm cross-entropy. Nghĩa là cực tiểu hóa cross-entropy "
  "chính là tìm tham số làm cho dữ liệu quan sát được có khả năng xảy ra cao nhất (ước lượng hợp lý cực đại).")
H2("4.3 Đạo hàm, gradient và quy tắc dây chuyền")
P("Để giảm mức phạt thì phải biết chỉnh trọng số theo hướng nào. Đạo hàm riêng của L theo một trọng số cho biết nếu tăng nhẹ trọng số đó thì mức phạt tăng hay giảm. "
  "Gom đạo hàm theo mọi trọng số lại được gradient, vector chỉ hướng dốc lên nhanh nhất; đi ngược gradient thì mất mát giảm. "
  "L phụ thuộc w qua một chuỗi: w → z → p → L, nên đạo hàm được tính bằng quy tắc dây chuyền, nhân các đạo hàm từng khâu.")
F("∂L/∂z = ∂L/∂p · ∂p/∂z = [ (p − y) / (p(1 − p)) ] · p(1 − p) = p − y")
F("∇_w L = (1/n) Xᵀ (p − y),        ∂L/∂b = (1/n) Σᵢ (pᵢ − yᵢ)")
P("Điều đẹp ở đây là các thừa số p(1 − p) triệt tiêu nhau, và kết quả cuối rất dễ đọc: gradient theo trọng số của đặc trưng j là trung bình của "
  "(sai số dự đoán) nhân (giá trị đặc trưng j). Nếu mô hình thường dự đoán thấp hơn thực tế (p nhỏ hơn y) ở những sinh viên có đặc trưng j lớn, "
  "thì trọng số đó sẽ được đẩy lên. Chính công thức này được cài thẳng trong code của nhóm.")
H2("4.4 Ảnh hưởng của các tham số đến mô hình")
w = dict(zip(cmpw.dac_trung, cmpw.w_tu_cai))
import math
names = [("Tuition fees up to date", "Đóng học phí đủ hạn"), ("Debtor", "Đang nợ"), ("Scholarship holder", "Có học bổng"),
         ("Age at enrollment", "Tuổi nhập học"), ("Curricular units 1st sem (approved)", "Số môn qua HK1"),
         ("Curricular units 1st sem (grade)", "Điểm HK1"), ("Curricular units 2nd sem (approved)", "Số môn qua HK2"),
         ("Curricular units 2nd sem (grade)", "Điểm HK2")]
rows = [[vn, f(w[k]), f(math.exp(w[k]), 2)] for k, vn in names]
TABLE(["Đặc trưng (đã chuẩn hóa)", "Trọng số w", "Tỉ số odds e^w"], rows, widths=[7, 3.5, 3.5])
P("Tỉ số odds e^w cho biết khi đặc trưng đó tăng một độ lệch chuẩn (các biến khác giữ nguyên) thì odds bỏ học nhân với bao nhiêu. "
  "Đóng đủ học phí có e^w ≈ 0,42, tức odds bỏ học giảm hơn một nửa; số môn qua HK2 có e^w ≈ 0,12, tác động mạnh nhất; có học bổng là 0,59. "
  "Còn nợ học phí (1,38) và tuổi lớn (1,35) đẩy nguy cơ lên. Tất cả đều khớp trực giác.")
P("Có một chi tiết lạ: trọng số của số môn qua HK1 và điểm HK1 lại dương (1,41 và 1,29), nghĩa là tưởng như qua nhiều môn HK1 thì nguy cơ tăng. "
  "Cần đọc kết quả này cho đúng. Số môn qua HK1 và HK2 tương quan tới 0,92 nên chúng gần như thông tin thừa của nhau; khi cả hai nằm trong mô hình, "
  "trọng số của từng cột chỉ có nghĩa trong sự so sánh với cột kia. Một cách đọc hợp lý (nhóm chưa kiểm chứng riêng) là mô hình coi HK2 làm tín hiệu chính, "
  "còn phần HK1 đóng vai hiệu chỉnh: một em qua nhiều môn ở HK1 mà sụt mạnh ở HK2 là dấu hiệu đáng lo hơn một em luôn yếu. "
  "Vì vậy không nên đọc từng trọng số tách rời khi các đặc trưng cộng tuyến; đây là một hạn chế của việc diễn giải, không phải lỗi của thuật toán.")
H2("4.5 Cực tiểu hóa mất mát chính là quá trình học")
P("“Học” ở đây không có gì huyền bí: đó là tìm w và b làm L nhỏ nhất. Với Hồi quy Logistic, L là hàm lồi theo tham số, nên chỉ có một đáy duy nhất, "
  "đi xuống theo gradient từ đâu cũng tới cùng một nơi. Đó là lý do mô hình này rất hợp để quan sát Gradient Descent. "
  "TabICLv2 cũng được huấn luyện bằng một dạng cross-entropy, nhưng trên mạng Transformer rất sâu, hàm mất mát không lồi và quy tắc dây chuyền "
  "phải đi qua nhiều lớp (lan truyền ngược). Điểm khác biệt quan trọng là việc huấn luyện đó chỉ diễn ra một lần, do tác giả làm trên dữ liệu tổng hợp; "
  "khi dùng cho sinh viên của ta, mô hình không cập nhật trọng số nữa.")

# ======================= 5 =======================
H1("5. Cơ sở toán học của thuật toán tối ưu")
BOX("Nói dễ hiểu:", '“Học” thực chất là chỉnh các trọng số cho mức phạt nhỏ dần. Gradient Descent làm việc đó như người đi xuống đồi trong sương mù: mỗi bước nhìn độ dốc dưới chân rồi bước xuống một đoạn. Đoạn bước (tốc độ học) quá lớn thì nhảy qua đáy, quá nhỏ thì đi mãi. Phần thí nghiệm sẽ cho thấy điều đó bằng số liệu thật.', fill="E2F0D9")
H2("5.1 Bài toán tối ưu tương ứng")
P("Với mô hình đối chứng, bài toán là cực tiểu hóa cross-entropy L(w, b) trên tập huấn luyện, không có ràng buộc (ràng buộc xác suất nằm trong [0, 1] "
  "đã được sigmoid đảm bảo). Vì L lồi và khả vi nên Gradient Descent là lựa chọn tự nhiên và đáng tin.")
H2("5.2 Gradient Descent và ý nghĩa công thức cập nhật")
F("w ← w − η · ∇_w L,        b ← b − η · ∂L/∂b")
P("Hãy hình dung đang đứng trên sườn đồi trong sương mù, chỉ cảm nhận được độ dốc dưới chân. Gradient là la bàn chỉ hướng dốc lên; ta bước ngược lại một đoạn. "
  "Đoạn bước dài bao nhiêu do tốc độ học η quyết định. Bước quá ngắn thì đi mãi không tới đáy; bước quá dài thì nhảy qua đáy, rơi sang sườn bên kia, "
  "thậm chí càng nhảy càng cao. Có thể chứng minh với hàm có gradient liên tục Lipschitz hằng số L_c thì η < 2/L_c là điều kiện để mỗi bước không làm hàm tăng; "
  "với cross-entropy logistic, hằng số này không vượt quá một phần tư trị riêng lớn nhất của XᵀX/n. Nhóm dùng chính điều này để tiên đoán tốc độ học an toàn ở mục 5.4.")
P("Nhóm dừng thuật toán khi một trong hai điều kiện đúng: độ giảm của mất mát giữa hai vòng liên tiếp nhỏ hơn 10⁻¹² hoặc đạt số vòng lặp tối đa (5.000 vòng khi so với sklearn, 2.000 vòng khi khảo sát tốc độ học). Khởi tạo w = 0 và b = 0.")
H2("5.3 Muon và cách TabICLv2 được tối ưu")
P("Theo bài gốc, TabICLv2 được tiền huấn luyện trên dữ liệu tổng hợp với mục tiêu cực tiểu hóa kỳ vọng của −log q_θ(y_test | x_test, D_train), "
  "tức cross-entropy của việc đoán nhãn các hàng test khi đã cho xem tập huấn luyện làm ngữ cảnh. Điểm đổi mới về tối ưu là các tác giả dùng bộ tối ưu Muon thay cho AdamW. "
  "Hiểu ở mức nguyên lý, Muon là bộ tối ưu có momentum, cập nhật các ma trận trọng số theo hướng đã được chỉnh để các hướng cập nhật có độ lớn đồng đều hơn "
  "(cần đối chiếu chi tiết với bài gốc và tài liệu Muon khi nộp bản cuối).")
BOX("Giới hạn của thí nghiệm tối ưu:",
    "vì TabICLv2 không huấn luyện lại trên dữ liệu sinh viên, nhóm không thể thay đổi tốc độ học của Muon. "
    "Do đó thí nghiệm về yếu tố tối ưu mà đề bài yêu cầu được thực hiện trên mô hình đối chứng (Hồi quy Logistic với Gradient Descent tự cài). "
    "Với TabICLv2, các yếu tố có thể thay đổi lúc suy luận là kích thước ngữ cảnh và số thành viên ensemble (mục 7.5).")
H2("5.4 Thực nghiệm: thay đổi tốc độ học và việc chuẩn hóa")
P("Nhóm huấn luyện Hồi quy Logistic bằng Gradient Descent tự cài trên tám đặc trưng, thử bốn tốc độ học (0,001; 0,01; 0,1; 1,0), "
  f"mỗi cái hai lần: có và không chuẩn hóa dữ liệu. Giá trị mất mát tối ưu (do sklearn tìm) trên tập huấn luyện là {f(info['sklearn_train_loss_optimum'],4)}.")
rows = []
for _, r in opt.iterrows():
    reach = "chưa tới" if pd.isna(r.vong_lap_den_gan_toi_uu) else str(int(r.vong_lap_den_gan_toi_uu))
    rows.append(["Có" if r.chuan_hoa else "Không", str(r.toc_do_hoc).replace(".", ","), f(r.loss_cuoi, 4), reach,
                 int(r.so_lan_loss_tang), pc(r.test_accuracy), pc(r.test_recall), f(r.test_auc, 3)])
TABLE(["Chuẩn hóa", "Tốc độ học η", "Loss sau 2.000 vòng", "Vòng lặp tới gần tối ưu*", "Số lần loss tăng", "Accuracy test", "Recall test", "AUC test"],
      rows, size=9, hl=[3])
P("* “Gần tối ưu” là mất mát nhỏ hơn mức tối ưu cộng 0,001. “Chưa tới” nghĩa là sau 2.000 vòng vẫn chưa đạt. Hàng tô vàng là cấu hình tốt nhất.", italic=True, size=10.5)
FIG("h5_duong_loss.png", "Hình 5. Đường mất mát theo số vòng lặp (trục ngang theo thang log). Vạch đứt là mức tối ưu.", 15.5)
P("Nhìn vào bảng và hình, có bốn nhận xét. Thứ nhất, khi đã chuẩn hóa, tốc độ học càng lớn thì càng nhanh: η = 1,0 chỉ cần 145 vòng để tới gần đáy, "
  "η = 0,1 cần 1.437 vòng, còn η = 0,01 và 0,001 sau 2.000 vòng vẫn chưa tới nơi. Mất mát luôn giảm đều (không lần nào tăng), đúng như lý thuyết cho hàm lồi. "
  "Thứ hai, khi không chuẩn hóa thì cùng một tốc độ học cho kết quả tệ hơn rất nhiều: η = 0,1 làm mất mát nhảy lên xuống (912 lần tăng) và kết thúc ở 1,043; "
  "η = 1,0 thì phân kỳ, mất mát lên tới 9,67 và Accuracy chỉ còn 73,7%.")
P(f"Thứ ba, lý do nằm ở số điều kiện đã tính ở mục 3.6. Không chuẩn hóa, trị riêng lớn nhất của XᵀX/n là khoảng {cond['khong_chuan_hoa']['lam_max']:.0f} (do cột tuổi và điểm có giá trị lớn), "
  f"nên tốc độ học an toàn tối đa 2/L_c chỉ khoảng {f(cond['khong_chuan_hoa']['lr_on_dinh_toi_da_2_over_L'],4)}. Vì thế η = 0,01 còn đi được, còn 0,1 và 1,0 vượt ngưỡng nên dao động hoặc phân kỳ. "
  f"Có chuẩn hóa, trị riêng lớn nhất chỉ khoảng {f(cond['chuan_hoa']['lam_max'],1)}, ngưỡng an toàn lên tới khoảng {f(cond['chuan_hoa']['lr_on_dinh_toi_da_2_over_L'],2)}, nên η = 1,0 vẫn ổn. "
  "Lý thuyết và thực nghiệm khớp nhau. Thứ tư, khi hội tụ, các cấu hình chuẩn hóa cho kết quả giống nhau về chất lượng (Accuracy quanh 90%), "
  "nghĩa là tốc độ học ảnh hưởng chủ yếu đến việc có tới đáy được hay không và nhanh đến đâu, chứ không đổi đáy ở đâu.")
H2("5.5 Đối chiếu bản tự cài với thư viện")
P(f"Đề bài khuyến khích tự xây dựng bước cập nhật để kiểm chứng. Nhóm chạy Gradient Descent tự cài (η = 1,0, dừng sau {th(info['gd_iters_to_converge'])} vòng) "
  f"và so với LogisticRegression của scikit-learn không điều chuẩn. Độ lệch lớn nhất giữa hai bộ trọng số chỉ là {f(info['max_abs_diff_weights'],4)}, "
  "nghĩa là hai bên cùng tìm ra một đáy. Điều này xác nhận công thức gradient ở mục 4.3 được cài đúng.")
rows = [[r.dac_trung.replace("Curricular units", "CU"), f(r.w_tu_cai, 4), f(r.w_sklearn, 4)] for _, r in cmpw.iterrows()]
TABLE(["Đặc trưng", "w (tự cài)", "w (sklearn)"], rows, widths=[8, 3.5, 3.5], size=9.5)

# ======================= 6 =======================
H1("6. Cơ sở xác suất và thống kê trong học máy")
BOX("Nói dễ hiểu:", 'Một mô hình đoán đúng trên bài đã học chưa chắc đoán đúng với sinh viên mới. Vì vậy ta giấu một phần dữ liệu để thi thử, và biết rằng điểm thi thử cũng có sai số ngẫu nhiên. Phần này nói cách chia dữ liệu, nhận ra học vẹt (overfitting), và biết chênh lệch bao nhiêu điểm thì mới đáng tin.', fill="E2F0D9")
H2("6.1 Những thành phần xác suất của bài toán")
P("Xác suất xuất hiện ngay từ đầu ra của mô hình: p = P(y = 1 | x) là xác suất có điều kiện của việc bỏ học khi biết hồ sơ x. Nhãn của mỗi sinh viên được coi như một "
  "biến ngẫu nhiên Bernoulli với tham số p; kỳ vọng của nó chính là p, và phương sai là p(1 − p). Với p = 0,5 độ bất định lớn nhất (phương sai 0,25), "
  "với p gần 0 hoặc 1 mô hình gần như chắc chắn. Việc ước lượng tham số là ước lượng hợp lý cực đại, đã nói ở mục 4.2: tìm w, b làm cho dữ liệu quan sát có khả năng cao nhất.")
P("Với TabICLv2, xác suất còn đi xa hơn một bước: câu trả lời là P(y | x, D_train), phụ thuộc cả vào tập huấn luyện được đưa vào ngữ cảnh. "
  "Đây là khung huấn luyện kiểu PFN (prior-data fitted network) mà bài gốc dùng: mô hình được huấn luyện trên dữ liệu sinh ra từ một phân phối giả định (prior), nên về bản chất nó xấp xỉ suy luận Bayes theo prior đó (cần đối chiếu bài gốc). "
  "Nói cách khác, nó đã “thấy” rất nhiều bài toán bảng giả lập; theo tài liệu của gói tabicl, nó được tiền huấn luyện cho các bảng khoảng 300 đến 48.000 dòng.")
H2("6.2 Chia dữ liệu huấn luyện và kiểm tra")
P("Học thuộc “bàn cũ” thì dễ; câu hỏi thật là gặp lớp sinh viên mới thì mô hình có đoán đúng không. Mất mát trên tập huấn luyện chỉ là ước lượng "
  "lạc quan của mất mát kỳ vọng trên dữ liệu chưa thấy, vì tham số đã được chỉnh để khớp đúng dữ liệu đó. Vì thế ta giấu đi một phần dữ liệu (tập test) "
  "và chỉ dùng nó để chấm điểm sau cùng.")
P(f"Nhóm chia 80/20 theo kiểu phân tầng, giữ nguyên tỉ lệ nhãn, với hạt giống ngẫu nhiên cố định (seed = 42): {th(info['n_train'])} sinh viên huấn luyện và {info['n_test']} sinh viên kiểm tra. "
  f"Tỉ lệ bỏ học ở hai tập lần lượt là {pc(info['dropout_rate_train'],2)} và {pc(info['dropout_rate_test'],2)}, gần như bằng nhau. "
  "Cách chia này được dùng chung cho mọi cấu hình, kể cả TabICLv2, để so sánh công bằng. Với TabICLv2, “huấn luyện” nghĩa là đưa các sinh viên này vào ngữ cảnh.")
H2("6.3 Overfitting, underfitting và khả năng tổng quát hóa")
P("Overfitting là khi mô hình nhớ cả nhiễu của dữ liệu huấn luyện nên điểm trên tập huấn luyện cao mà trên tập mới thấp; underfitting là khi mô hình quá đơn giản "
  "nên điểm thấp ở cả hai. Hồi quy Logistic chỉ có chín tham số cho gần ba nghìn mẫu nên rất khó quá khớp, và số liệu xác nhận điều này: "
  f"Accuracy huấn luyện {pc(cfg.loc[2,'train_acc'])} và kiểm tra {pc(cfg.loc[2,'accuracy'])} (cấu hình B), không có khoảng cách đáng ngại. "
  f"Hơn nữa mất mát kiểm tra ({f(cfg.loc[2,'test_loss'])}) còn thấp hơn mất mát huấn luyện ({f(cfg.loc[2,'train_loss'])}); đây là do tập test ngẫu nhiên hơi dễ hơn, không phải dấu hiệu bất thường.")
P("Mô hình attention minh họa cho thấy cả hai thái cực ở mục 7.4. Tham số τ điều khiển việc “chú ý” rộng hay hẹp. "
  "Khi τ rất nhỏ, mô hình chỉ nhìn vài người giống nhất, quá nhạy với nhiễu cục bộ (giống nhớ vặt). Khi τ rất lớn, mô hình nhìn mọi người gần như đều nhau "
  "và biến thành dự đoán theo trung bình chung, bỏ sót nhiều em bỏ học (underfitting). Hai thái cực này là ví dụ trực quan cho thế cân bằng giữa độ chệch và phương sai.")
P("Một lưu ý về độ tin cậy: với 726 mẫu kiểm tra, sai số chuẩn của Accuracy quanh mức 90% là khoảng 1,1 điểm phần trăm, tức khoảng tin cậy 95% xấp xỉ ±2,2 điểm. "
  "Vì vậy chênh lệch nhỏ hơn 2 điểm giữa hai cấu hình không đủ để kết luận cấu hình nào tốt hơn; nhóm chỉ kết luận những chênh lệch lớn hơn rõ rệt. "
  "Báo cáo này dùng một lần chia duy nhất, nên chưa có thông tin về độ ổn định giữa các lần chia (kiểm tra chéo k-fold là hướng mở rộng).")

# ======================= 7 =======================
H1("7. Xây dựng mô hình và thực nghiệm")
BOX("Nói dễ hiểu:", 'Đây là phần “chấm điểm”. Ta thử nhiều cấu hình (đưa ít hay nhiều thông tin vào mô hình) và đo bằng các độ đo phù hợp. Trong bài này, bắt được các em sắp bỏ học (Recall) quan trọng hơn là tỉ lệ đúng chung, vì bỏ sót một em thì mất một sinh viên.', fill="E2F0D9")
H2("7.1 Tiền xử lý và cơ sở toán học")
BUL([("Bỏ nhóm “đang học”: ", "vì chưa có kết cục (mục 1.2)."),
     ("Chọn đặc trưng: ", "tám cột ở mục 1.3 cho cấu hình chính; hai cấu hình khác để so sánh (mục 7.3)."),
     ("Chuẩn hóa z-score ", "(chỉ dùng cho mô hình đối chứng và attention minh họa; trung bình và độ lệch chuẩn tính trên tập huấn luyện rồi áp lên tập test để tránh rò rỉ thông tin). Cơ sở: số điều kiện, mục 3.6 và 5.4."),
     ("Chia train/test phân tầng 80/20, seed 42: ", "mục 6.2."),
     ("Với TabICLv2: ", "theo tài liệu của gói tabicl, cột dạng phân loại và giá trị thiếu được gói tự xử lý, nên script đưa dữ liệu thô vào (cần đối chiếu tài liệu khi chạy).")])
H2("7.2 Các độ đo và ý nghĩa")
P("Mọi độ đo đều xuất phát từ bốn con số của ma trận nhầm lẫn: TP (đoán bỏ học và đúng là bỏ học), FN (bỏ sót, thật sự bỏ học mà đoán an toàn), "
  "FP (báo nhầm) và TN (đoán an toàn và đúng).")
TABLE(["Độ đo", "Công thức", "Ý nghĩa trong bài này"],
      [["Accuracy", "(TP + TN) / n", "Tỉ lệ đoán đúng chung; dễ đánh lừa khi lớp lệch"],
       ["Precision", "TP / (TP + FP)", "Trong số em bị báo nguy cơ, bao nhiêu em thật sự bỏ học (đỡ báo nhầm)"],
       ["Recall", "TP / (TP + FN)", "Trong số em thật sự bỏ học, bắt được bao nhiêu (đỡ bỏ sót em có nguy cơ)"],
       ["F1", "2·P·R / (P + R)", "Trung bình điều hòa của Precision và Recall"],
       ["ROC-AUC", "Diện tích dưới đường ROC", "Xác suất mô hình chấm một em bỏ học cao hơn một em an toàn, không phụ thuộc ngưỡng"]],
      widths=[2.6, 4.5, 9.0], size=10)
P("Hai kiểu sai có giá rất khác nhau. Báo nhầm một em ổn là nguy cơ chỉ tốn một buổi gặp mặt; bỏ sót một em sắp nghỉ học là em đó “nổ” thật. "
  "Vì vậy Recall của lớp bỏ học là độ đo quan trọng nhất ở bài này, và Accuracy chỉ đóng vai phụ.")
H2("7.3 Các cấu hình thực nghiệm với mô hình đối chứng")
P("Để phân tích ảnh hưởng của một thành phần đến hiệu quả, nhóm thay đổi lượng thông tin đưa vào mô hình, từ thông tin lúc nhập học, "
  "đến hết học kỳ 1, hết học kỳ 2, và toàn bộ 35 cột. Cùng một cách chia dữ liệu, cùng thuật toán (Hồi quy Logistic đã chuẩn hóa).")
rows = [[r.cau_hinh.replace("Luc nhap hoc", "Lúc nhập học").replace("Nhap hoc + het", "Nhập học + hết").replace("Day du 35 cot", "Đầy đủ 35 cột").replace("cot gon", "cột gọn").replace("cot", "cột"),
         pc(r.accuracy), pc(r.precision), pc(r.recall), pc(r.f1), f(r.auc)] for _, r in cfg.iterrows()]
TABLE(["Cấu hình", "Accuracy", "Precision", "Recall (bỏ học)", "F1", "AUC"], rows, widths=[6.4, 2, 2, 2.4, 1.8, 1.8], size=9.5, hl=[2])
P("Kết quả cho thấy rõ thông tin càng nhiều thì dự đoán càng tốt, nhưng cũng đi kèm cái giá về thời điểm. Chỉ dùng thông tin lúc nhập học, Recall mới đạt 58,1%: "
  "gần một nửa số em bỏ học bị bỏ sót, vì lúc đó mô hình chưa biết em học thế nào. Thêm kết quả học kỳ 1 đẩy Recall lên 80,6% và AUC từ 0,815 lên 0,914. "
  "Thêm học kỳ 2 lên 85,6% (Accuracy 89,9%). Dùng đủ 35 cột cho kết quả tốt nhất (Accuracy 94,5%, Recall 92,3%, AUC 0,973). "
  "Chênh lệch giữa A và các cấu hình sau là rất lớn so với sai số ±2,2 điểm nên đáng tin; chênh lệch giữa B1 và B (2,5 điểm Accuracy) thì ở ngưỡng biên.")
P("Cũng cần nhắc lại điều ở mục 2.5: bản “hết học kỳ 2” có phần được hưởng lợi vì điểm 0 chính là dấu hiệu em đã nghỉ. "
  "Nếu mục tiêu là cảnh báo sớm thật sự thì cấu hình B1 (hết học kỳ 1) mới là phương án thực tế, với Recall 80,6%.")
FIG("h7_nhamlan.png", "Hình 6. Ma trận nhầm lẫn của cấu hình B ở ngưỡng 0,5 trên 726 sinh viên kiểm tra.", 6.5)
P("Ngưỡng 0,5 không bắt buộc. Nếu nhà trường muốn phát hiện nhiều em có nguy cơ hơn, có thể hạ ngưỡng báo động xuống. Bảng dưới (cấu hình B) cho thấy đây là một đánh đổi: "
  "ngưỡng càng thấp thì Recall càng cao nhưng Precision và Accuracy càng giảm.")
rows = [[f(r.nguong, 1), pc(r.recall), pc(r.precision), pc(r.accuracy), int(r.TP), int(r.FN), int(r.FP)] for _, r in thr.iterrows()]
TABLE(["Ngưỡng", "Recall", "Precision", "Accuracy", "TP", "FN (bỏ sót)", "FP (báo nhầm)"], rows, size=9.5, hl=[1])
P("Ở ngưỡng 0,3, mô hình bắt được 91,5% em bỏ học (chỉ bỏ sót 24 em) đổi lại báo nhầm 71 em, trong khi ở ngưỡng 0,5 bỏ sót 41 em và báo nhầm 32 em. "
  "Nếu mỗi lần báo nhầm chỉ tốn một buổi gặp mặt còn mỗi lần bỏ sót là mất một sinh viên, thì ngưỡng thấp hơn 0,5 là lựa chọn hợp lý. "
  "Đây là quyết định của nhà trường, không phải của thuật toán; AUC không đổi (0,940) vì thứ hạng xác suất không đổi.")
FIG("h6_nguong.png", "Hình 7. Đánh đổi giữa Recall, Precision và Accuracy khi đổi ngưỡng (cấu hình B).", 9.5)
H2("7.4 Mô hình attention minh họa: học trong ngữ cảnh trông như thế nào")
P("Để quan sát cơ chế “học trong ngữ cảnh” mà không cần tải TabICLv2, nhóm tự viết một đầu attention nhỏ, không có trọng số học được: "
  "sinh viên cần dự đoán so sánh với các sinh viên trong ngữ cảnh bằng khoảng cách (như ví dụ ở mục 3.5), softmax biến khoảng cách thành trọng số, "
  "và xác suất bỏ học là tổng trọng số của những người đã bỏ học. Tham số τ điều khiển độ rộng của sự chú ý. "
  "Đây chỉ là bản minh họa cùng ý tưởng cơ bản, hoàn toàn không phải TabICLv2 và không nên dùng để suy ra chất lượng của TabICLv2.")
sel = att[att.tau.isin([0.05, 0.5, 8.0])]
rows = [[int(r.so_mau_ngu_canh), str(r.tau).replace(".", ","), pc(r.accuracy), pc(r.recall), f(r.auc)] for _, r in sel.iterrows()]
TABLE(["Số sinh viên trong ngữ cảnh", "τ", "Accuracy", "Recall (bỏ học)", "AUC"], rows, size=9.5)
FIG("h8_attention.png", "Hình 8. AUC của attention minh họa theo kích thước ngữ cảnh, với bốn giá trị τ.", 10.5)
P("Có hai quy luật đáng chú ý. Ngữ cảnh càng nhiều thì AUC càng cao (với τ = 0,5: từ 0,901 ở 50 sinh viên lên 0,947 ở toàn bộ 2.904 sinh viên), vì có thêm nhiều ví dụ giống để tham khảo. "
  "Ngoài ra, τ vừa phải (0,5) cho kết quả tốt nhất; τ rất nhỏ (0,05) nhảy theo vài người giống nhất (AUC 0,918 ở ngữ cảnh đầy đủ, thấp hơn 0,947), "
  "còn τ lớn (8,0) làm Recall tụt xuống 42,3% vì mọi sinh viên trông gần như “giống nhau” với tất cả. Chính cơ chế chọn mức tập trung cho attention, "
  "kết hợp với việc học nó từ rất nhiều bài toán giả lập, là điều TabICLv2 làm bằng tay một cách tinh vi hơn nhiều.")
H2("7.5 Kết quả TabICLv2 và so sánh với mô hình đối chứng")
P("Nhóm chạy TabICLv2 bằng lớp TabICLClassifier của thư viện tabicl (phiên bản 2.2.0) với trọng số tabicl-classifier-v2-20260212 và 8 thành viên ensemble (mặc định). "
  "Dữ liệu thô được đưa vào, không chuẩn hóa tay vì mô hình tự xử lý thang đo. Toàn bộ 2.904 sinh viên huấn luyện đóng vai ngữ cảnh, "
  "726 sinh viên kiểm tra là những em cần dự đoán, cùng cách chia (seed 42) và cùng bốn cấu hình như ở mục 7.3. Hàm fit chỉ đưa ngữ cảnh vào, không cập nhật trọng số nào.")
TABLE(["Cấu hình", "Mô hình", "Accuracy", "Precision", "Recall (bỏ học)", "F1", "AUC"],
      [["A (lúc nhập học)", "Hồi quy Logistic", "75,1%", "72,7%", "58,1%", "0,646", "0,815"],
       ["", "TabICLv2", "75,6%", "71,5%", "62,7%", "0,668", "0,823"],
       ["B1 (hết HK1)", "Hồi quy Logistic", "87,5%", "86,4%", "80,6%", "0,834", "0,914"],
       ["", "TabICLv2", "87,9%", "86,3%", "82,0%", "0,841", "0,934"],
       ["B (hết HK2)", "Hồi quy Logistic", "89,9%", "88,4%", "85,6%", "0,869", "0,940"],
       ["", "TabICLv2", "91,3%", "90,8%", "86,6%", "0,886", "0,955"],
       ["C (35 cột)", "Hồi quy Logistic", "94,5%", "93,6%", "92,3%", "0,929", "0,973"],
       ["", "TabICLv2", "94,6%", "94,2%", "91,9%", "0,930", "0,977"]],
      widths=[3.0, 3.4, 2.0, 2.0, 2.4, 1.6, 1.6], size=9.5, hl=[1, 3, 5, 7])
P("Thời gian dự đoán của TabICLv2 chỉ vài giây cho mỗi cấu hình (từ 1,1 đến 7,0 giây, lần đầu lâu hơn do khởi tạo), vì chỉ là một lượt tính, không có vòng lặp huấn luyện.")
P("Đọc kết quả. TabICLv2 cho AUC cao hơn mô hình đối chứng ở cả bốn cấu hình (hơn 0,008; 0,020; 0,015 và 0,004), và Recall cao hơn ở ba cấu hình đầu "
  "(hơn 4,6; 1,4 và 1,0 điểm phần trăm), trong khi ở cấu hình C hai mô hình gần như ngang nhau (Recall thấp hơn 0,4 điểm). "
  "Chênh lệch Accuracy nằm trong khoảng 0,1 đến 1,4 điểm, đều nhỏ hơn sai số khoảng ±2,2 điểm của phép đo (mục 6.3), nên chưa đủ để nói TabICLv2 chính xác hơn rõ rệt. "
  "Điều nhất quán hơn cả là AUC: mô hình xếp hạng nguy cơ tốt hơn một chút, nhưng đây vẫn chỉ là một lần chia dữ liệu, chưa có kiểm tra chéo để khẳng định.")
P("Một dự đoán của nhóm trước khi chạy là lợi thế của mô hình phi tuyến sẽ rõ nhất ở cấu hình nhiều cột (C). Kết quả không ủng hộ dự đoán đó: "
  "lợi thế lớn nhất lại ở cấu hình ít thông tin (A và B1), còn ở C hai mô hình gần như bằng nhau. Nhóm chưa kiểm chứng nguyên nhân; "
  "một cách giải thích hợp lý (chưa được kiểm tra) là khi đã có đủ 35 cột thì mô hình tuyến tính đã khá mạnh, còn khi ít cột, khả năng so sánh với những sinh viên giống mình "
  "của attention giúp bù lại phần quan hệ phi tuyến mà mô hình tuyến tính bỏ lỡ.")
BOX("Phạm vi của kết quả:", "mới chạy một lần với cấu hình mặc định. Chưa khảo sát ảnh hưởng của kích thước ngữ cảnh và số thành viên ensemble, "
    "và chưa lặp lại trên nhiều lần chia dữ liệu, nên không nên suy ra điều gì mạnh hơn những gì bảng trên cho thấy.", fill="FFF2CC")

# ======================= 8 =======================
H1("8. Thảo luận và kết luận")
BOX("Nói dễ hiểu:", 'Phần cuối nối kết quả với toán: vì sao phải chuẩn hóa (số điều kiện), vì sao đổi ngưỡng thì Recall thay đổi (xác suất và chi phí sai), và những chỗ kết quả chưa đáng tin (TabICLv2 mới chạy một lần, điểm học kỳ 2 chứa sẵn đáp án). Nói thẳng hạn chế cũng là một phần của bài làm tốt.', fill="E2F0D9")
H2("8.1 Tổng hợp các kiến thức toán học quan trọng")
TABLE(["Mảng toán", "Dùng ở đâu trong bài", "Bằng chứng từ thực nghiệm"],
      [["Đại số tuyến tính", "Vector/ma trận biểu diễn sinh viên; tích vô hướng wᵀx; Q, K, V và attention; trị riêng, số điều kiện",
        f"Ví dụ tay attention (p ≈ {f(ex['p_bo_hoc'],3)}); số điều kiện {cond['chuan_hoa']['so_dieu_kien']:.0f} so với {th(cond['khong_chuan_hoa']['so_dieu_kien'])}"],
       ["Giải tích", "Sigmoid, cross-entropy, gradient Xᵀ(p − y)/n, quy tắc dây chuyền", "Bản tự cài khớp sklearn (lệch tối đa 0,0005)"],
       ["Tối ưu", "Gradient Descent, tốc độ học, điều kiện dừng; Muon (ở TabICLv2)", "Bảng 4 tốc độ học × có/không chuẩn hóa; η = 1,0 hội tụ sau 145 vòng"],
       ["Xác suất, thống kê", "P(y|x), MLE, chia train/test, bias–variance, sai số chuẩn", "Tỉ lệ nhãn hai tập 39,15% và 39,12%; τ nhỏ/lớn cho hai kiểu sai"],
       ["Đánh giá", "Ma trận nhầm lẫn, Precision/Recall/F1/AUC, đổi ngưỡng", "Ngưỡng 0,3 bắt 91,5% em bỏ học so với 85,6% ở ngưỡng 0,5"]],
      widths=[3, 7, 6], size=9.5)
H2("8.2 Giải thích kết quả bằng nguyên lý toán học")
P("Ba kết quả nổi bật của thực nghiệm đều giải thích được bằng toán đã trình bày. Một, Gradient Descent trên dữ liệu không chuẩn hóa dao động hoặc phân kỳ ở tốc độ học lớn, "
  "vì ma trận XᵀX/n có số điều kiện cao và ngưỡng η < 2/L_c nhỏ hơn rất nhiều; chuẩn hóa là phép biến đổi đại số tuyến tính đơn giản mà thay đổi hoàn toàn khả năng hội tụ. "
  "Hai, thông tin càng nhiều thì Recall càng cao nhưng phải trả giá bằng thời điểm cảnh báo, và đó là sự đánh đổi giữa lượng thông tin và tính kịp thời, "
  "không phải sự đánh đổi của thuật toán. Ba, mô hình cân nhắc giữa bỏ sót và báo nhầm thông qua ngưỡng, và cross-entropy phạt nặng sự tự tin sai là lý do mô hình cho xác suất đáng tin để chọn ngưỡng.")
H2("8.3 Hạn chế và các giả định")
BUL([("TabICLv2 mới chạy một lần. ", "Chỉ một cấu hình mặc định trên một lần chia dữ liệu; chưa khảo sát kích thước ngữ cảnh, số thành viên ensemble, và thí nghiệm tốc độ học vẫn làm trên mô hình đối chứng vì TabICLv2 không huấn luyện lại."),
     ("Rò rỉ thông tin ở cấu hình dùng học kỳ 2. ", "Điểm 0 ở HK2 thường nghĩa là em đã nghỉ, nên cấu hình B và C có phần dự đoán một điều đã xảy ra."),
     ("Cộng tuyến. ", "Số môn qua HK1 và HK2 tương quan 0,92 nên trọng số riêng lẻ khó diễn giải (mục 4.4)."),
     ("Một lần chia dữ liệu, 726 mẫu kiểm tra. ", "Sai số khoảng ±2,2 điểm phần trăm; chưa dùng kiểm tra chéo."),
     ("Giả định độc lập và tuyến tính. ", "Hồi quy Logistic giả định mức nguy hiểm là tổ hợp tuyến tính của đặc trưng và các sinh viên độc lập với nhau."),
     ("Tính tổng quát. ", "Dữ liệu chỉ từ một trường ở Bồ Đào Nha, kết luận có thể không đúng ở nơi khác."),
     ("Đạo đức. ", "Dự đoán nguy cơ chỉ nên dùng để hỗ trợ sinh viên, không để quy kết; các biến nhạy cảm như giới tính cần cân nhắc kỹ.")])
H2("8.4 Kết luận")
P("Trả lời bốn câu hỏi định hướng của đề bài: bài toán được mô hình hóa thành phân loại nhị phân với đầu ra là xác suất bỏ học, dữ liệu là vector tám chiều của mỗi sinh viên. "
  "Cách thức hoạt động của thuật toán do bốn mảng toán quyết định: đại số tuyến tính để biểu diễn và gộp thông tin, giải tích để đo sai lệch và chỉ hướng sửa, "
  "tối ưu để thực sự tìm ra tham số, và xác suất thống kê để biết mô hình có tin được trên dữ liệu mới không. Tham số được học bằng cách cực tiểu hóa cross-entropy, "
  "với mô hình đối chứng bằng Gradient Descent (đã kiểm chứng khớp thư viện), còn TabICLv2 được tiền huấn luyện một lần bằng Muon rồi dùng không cần cập nhật; khi chạy thật trên cùng cách chia dữ liệu, TabICLv2 cho AUC cao hơn mô hình đối chứng ở cả bốn cấu hình nhưng chênh Accuracy nhỏ hơn sai số phép đo (mục 7.5). "
  "Kết quả thực nghiệm giải thích được bằng toán: số điều kiện giải thích vì sao phải chuẩn hóa, cross-entropy và ngưỡng giải thích cách cân giữa bỏ sót và báo nhầm.")
P("Nhóm rút ra một điều: đoán đúng hay sai không đáng giá bằng việc biết cái giá của từng kiểu sai. Toán học cho ta công cụ để đo cái giá đó, "
  "và việc chọn mô hình cho tốt chỉ là một nửa; nửa còn lại là chọn đúng thời điểm và đúng ngưỡng để hành động.")

# ======================= TAI LIEU =======================
H1("Tài liệu tham khảo")
for t in [
    "Realinho, V., Machado, J., Baptista, L., Martins, M.V. (2022). Predicting Student Dropout and Academic Success. Data (MDPI). Bộ dữ liệu: UCI Machine Learning Repository, DOI 10.24432/C5MC89. https://archive.ics.uci.edu/dataset/697/predict+students+dropout+and+academic+success",
    "Qu, J., Holzmüller, D., Varoquaux, G., Le Morvan, M. (2026). TabICLv2: A Better, Faster, Scalable, and Open Tabular Foundation Model. Proceedings of the 43rd ICML, PMLR 306:102477–102555. https://proceedings.mlr.press/v306/qu26a.html ; arXiv:2602.11139.",
    "Mã nguồn TabICL: https://github.com/soda-inria/tabicl ; bản cài đặt giáo dục: https://github.com/soda-inria/nanotabicl",
    "Vaswani, A. và cộng sự (2017). Attention Is All You Need. NeurIPS."]:
    p = doc.add_paragraph(style="List Number"); p.add_run(t); p.paragraph_format.space_after = Pt(3)

H1("Phụ lục A. TabICLv2 từ con số 0 (dành cho người chưa học học sâu)")
BOX("Lưu ý:", "phụ lục này giải thích ý tưởng bằng lời và ví dụ, tóm tắt từ bài gốc của TabICLv2; các chi tiết kỹ thuật cần đối chiếu với bài báo khi nộp bản cuối. Kết quả khi nhóm chạy TabICLv2 trên dữ liệu sinh viên nằm ở mục 7.5.", fill="FCE4D6")
H2("A.1 Hai cách “học” khác nhau")
P("Hồi quy Logistic học theo kiểu thông thường: ta đưa dữ liệu của mình vào, máy chỉnh trọng số nhiều vòng cho tới khi mức phạt nhỏ. Dữ liệu của mình là thứ duy nhất nó học.")
P("TabICLv2 thì khác. Nó đã được học xong từ trước, một lần duy nhất, trên một lượng rất lớn bảng dữ liệu do máy tự sinh ra. Khi ta dùng, ta chỉ đưa cho nó danh sách các sinh viên đã biết kết cục (gọi là ngữ cảnh) kèm sinh viên cần đoán, và nó trả lời ngay trong một lượt tính, không chỉnh trọng số nào. Giống một thầy giáo đã dạy hàng chục nghìn lớp: gặp lớp mới, chỉ cần nhìn qua hồ sơ các em cũ là đoán được em mới. Cách này gọi là học trong ngữ cảnh (in-context learning).")
H2("A.2 Attention: so sánh rồi bỏ phiếu có trọng số")
P("Cốt lõi của mô hình là cơ chế attention. Có ba bước. Một, so sánh em cần đoán với từng em cũ để ra một điểm “giống nhau”. Hai, dùng hàm softmax đổi các điểm đó thành những trọng số dương cộng lại bằng 1 (càng giống, trọng số càng lớn). Ba, bỏ phiếu: xác suất bỏ học của em mới là tổng trọng số của những em cũ đã bỏ học.")
P("Ví dụ bằng số. Giả sử điểm giống nhau của em mới với bốn em cũ là 2, 1, 0, 0. Softmax tính e² ≈ 7,39; e¹ ≈ 2,72; e⁰ = 1; e⁰ = 1, tổng ≈ 12,11, nên trọng số là 0,61; 0,22; 0,08; 0,08. Nếu hai em đầu đã bỏ học, hai em sau thì không, xác suất bỏ học của em mới là 0,61 + 0,22 = 0,83.")
P("Trong mô hình thật, mỗi em được biến thành ba vector bằng ba ma trận trọng số học được: truy vấn Q (“tôi đang tìm ai giống tôi?”), khóa K (“tôi trông như thế nào?”) và giá trị V (“tôi mang thông tin gì?”). Điểm giống nhau là tích vô hướng giữa Q của em mới và K của em cũ, và công thức đầy đủ là softmax(QKᵀ/√d)·V. Phép chia cho √d chỉ để các con số không quá lớn làm softmax bị cứng. Bài gốc nói thông tin nhãn của em cũ được cộng vào biểu diễn của em đó, nhờ vậy kết cục đi theo trọng số sang em mới.")
H2("A.3 Vì sao có ba tầng")
BUL([("Tầng 1, nhìn từng cột: ", "mỗi chỉ số (học phí, tuổi, điểm...) có phân phối riêng, nên mô hình trước hết học cách biểu diễn từng cột trong cả bảng."),
     ("Tầng 2, nhìn từng hàng: ", "gộp các chỉ số của một sinh viên thành một vector đại diện duy nhất (như gộp tám con số thành một “hồ sơ”)."),
     ("Tầng 3, nhìn cả bảng: ", "các em cần đoán dùng attention để chú ý vào các em đã có nhãn, và cho ra xác suất.")])
P("Cách chia này giúp mô hình xử lý được bảng có số cột và số hàng khác nhau, vì mỗi tầng chỉ lo một loại quan hệ.")
H2("A.4 Học trước trên dữ liệu giả lập, vì sao vẫn dùng được cho dữ liệu thật")
P("Các tác giả không dùng dữ liệu thật để học trước, mà sinh ra rất nhiều bảng giả từ các “mô hình nhân quả ngẫu nhiên”: chọn ngẫu nhiên cái nào ảnh hưởng cái nào, bằng hàm gì, rồi sinh số liệu. Mô hình được luyện để đoán nhãn của các hàng kiểm tra khi đã cho xem các hàng huấn luyện của cùng bảng. Vì phải làm tốt trên vô số bảng khác nhau, nó học được một “kinh nghiệm chung” về cách quan hệ giữa các cột và nhãn thường xuất hiện. Hy vọng, và bài gốc cho thấy bằng thực nghiệm, là bảng thật cũng giống một trong các kiểu đó. Đây là điểm có thể bị hỏi: nếu dữ liệu thật lạ hẳn so với những gì đã học thì kết quả có thể kém.")
H2("A.5 Hai cải tiến kỹ thuật được nhắc tới")
BUL([("QASSMax: ", "khi ngữ cảnh có hàng nghìn em, softmax thường làm sự chú ý bị “loãng” vì mẫu số quá lớn. QASSMax điều chỉnh điểm số theo kích thước ngữ cảnh để mô hình vẫn tập trung được."),
     ("Bộ tối ưu Muon: ", "thay vì AdamW, các tác giả dùng Muon khi huấn luyện trước. Muon chỉnh hướng cập nhật của các ma trận trọng số bằng một phép lặp (Newton–Schulz) để các hướng có độ lớn đồng đều hơn. Nó chỉ dùng lúc huấn luyện trước, nên ta không thể chỉnh nó trên dữ liệu sinh viên.")])
H2("A.6 Khác biệt cần nhớ khi trình bày")
TABLE(["", "Hồi quy Logistic (đối chứng)", "TabICLv2"],
      [["Cách học", "Gradient Descent trên dữ liệu của mình", "Học trước một lần trên dữ liệu giả lập, không cập nhật khi dùng"],
       ["Số tham số", "9 (8 trọng số và 1 hệ số chặn)", "Khoảng 28 triệu (theo bài gốc)"],
       ["Có nhìn thấy toán không", "Từng bước, có thể tự tính tay", "Chỉ hiểu được ý tưởng, không tính tay trọn vẹn"],
       ["Trạng thái trong bài này", "Đã chạy, có số liệu", "Đã chạy (notebook 05), bốn cấu hình, một lần chia dữ liệu"]],
      widths=[3.6, 5.8, 6.6], size=10)
H1("Phụ lục B. Bảng thuật ngữ")
TABLE(["Thuật ngữ", "Hiểu đơn giản"],
      [["Vector, ma trận", "Một dãy số; một bảng số (mỗi hàng một sinh viên)."],
       ["Tích vô hướng", "Nhân từng cặp số rồi cộng lại; dùng để gộp chỉ số theo trọng số hoặc đo độ giống."],
       ["Sigmoid", "Hàm đổi số bất kỳ thành xác suất từ 0 đến 1."],
       ["Cross-entropy", "Thước đo mức phạt: đoán sai mà tự tin thì phạt rất nặng."],
       ["Ước lượng hợp lý cực đại (MLE)", "Chọn tham số sao cho dữ liệu thật có khả năng xảy ra cao nhất; với bài này tương đương cực tiểu hóa cross-entropy."],
       ["Gradient", "Vector chỉ chiều làm hàm tăng nhanh nhất; đi ngược lại thì hàm giảm."],
       ["Tốc độ học", "Độ dài mỗi bước khi đi xuống trong Gradient Descent."],
       ["Hàm lồi", "Hàm chỉ có một đáy, đi xuống từ đâu cũng tới cùng chỗ."],
       ["Trị riêng, số điều kiện", "Trị riêng cho biết độ dốc theo từng hướng; số điều kiện là tỉ số lớn nhất trên nhỏ nhất, càng lớn thì bề mặt càng méo và càng khó đi xuống."],
       ["Overfitting", "Học vẹt: điểm trên bài đã học cao nhưng trên bài mới thấp."],
       ["Recall, Precision", "Recall: trong số em thật sự bỏ học, bắt được bao nhiêu. Precision: trong số em bị báo nguy cơ, bao nhiêu em thật sự bỏ học."],
       ["AUC", "Xác suất mô hình chấm em bỏ học cao hơn em an toàn; không phụ thuộc ngưỡng."],
       ["Softmax", "Đổi một dãy điểm thành trọng số dương cộng lại bằng 1."],
       ["Attention", "So sánh một mục với các mục khác, đổi độ giống thành trọng số rồi lấy trung bình có trọng số."],
       ["Học trong ngữ cảnh", "Mô hình đã học sẵn; khi dùng chỉ cần đưa ví dụ có nhãn làm ngữ cảnh, không cập nhật trọng số."],
       ["Ensemble", "Chạy nhiều biến thể rồi lấy trung bình kết quả (TabICL mặc định 8 biến thể)."]],
      widths=[4.6, 11.4], size=10)
H1("Phụ lục C. Cách chạy lại")
BUL(["src/phan_tich.py: thống kê, biểu đồ, Hồi quy Logistic tự cài, attention minh họa; ghi mọi bảng và hình vào results/.",
     "notebooks/05_attention_va_tabicl.ipynb: ô cuối chạy TabICLv2 (cần internet để tải trọng số). src/chay_tabicl.py: bản chạy bằng dòng lệnh, chưa được chạy thử.",
     "src/viet_bao_cao.py: dựng lại báo cáo này từ các tệp trong results/.",
     "Hạt giống ngẫu nhiên: 42. Cách chia dữ liệu lưu trong results/split_seed42.json."])

doc.save(OUT)
print("saved", OUT)
